#!/usr/bin/env python3
"""Refresh the "Claude Code Status" column on the OTBI Subject Area Listing workbook.

The register lives in Google Drive as an .xlsx ("OTBI Subject Area Listing.xlsx").
This script adds, or refreshes, one column at the end of the "Subject Areas" sheet
saying what has been banked into otbi_meta. Nothing else in the workbook is touched.

The banked state comes from a JSON file, so the script stays offline and the
numbers come from a real query rather than being typed by hand. Produce it with:

    select json_object_agg(sa.subject_area, json_build_object(
             'folders', sa.folder_count,
             'columns', (select count(*) from otbi_meta.otbi_columns c
                          where c.subject_area = sa.subject_area),
             'formats', (select string_agg(distinct case c.source
                             when 'advanced_xml' then 'XML'
                             when 'advanced_sql' then 'SQL' else c.source end, ' + ')
                          from otbi_meta.otbi_columns c where c.subject_area = sa.subject_area),
             'date', to_char(sa.updated_at at time zone 'UTC','YYYY-MM-DD')))
    from otbi_meta.otbi_subject_areas sa where sa.capture_status = 'Captured';

Usage:
    python3 update_tracker.py LISTING.xlsx STATUS.json [-o OUT.xlsx]
"""

import argparse
import json
import sys
from copy import copy

import openpyxl
from openpyxl.utils import get_column_letter

SHEET = 'Subject Areas'
HEADER = 'Claude Code Status'
SUBJECT_AREA_HEADER = 'Subject Area'
NOT_BANKED = 'Not banked'


def find_column(ws, header):
    for cell in ws[1]:
        if cell.value == header:
            return cell.column
    return None


def refresh(path, statuses, out_path):
    wb = openpyxl.load_workbook(path)
    if SHEET not in wb.sheetnames:
        raise SystemExit('%s: no %r sheet, found %s' % (path, SHEET, wb.sheetnames))
    ws = wb[SHEET]

    sa_col = find_column(ws, SUBJECT_AREA_HEADER)
    if sa_col is None:
        raise SystemExit('%s: no %r column in the header row' % (path, SUBJECT_AREA_HEADER))

    col = find_column(ws, HEADER)
    if col is None:
        # New column at the end, styled like the existing Status column so it
        # does not look bolted on.
        col = ws.max_column + 1
        template = ws.cell(row=1, column=find_column(ws, 'Status') or 1)
        cell = ws.cell(row=1, column=col, value=HEADER)
        for attr in ('font', 'fill', 'border', 'alignment'):
            setattr(cell, attr, copy(getattr(template, attr)))
        ws.column_dimensions[get_column_letter(col)].width = 62

    body = ws.cell(row=2, column=find_column(ws, 'Status') or 1)
    seen, banked = set(), 0
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row=row, column=sa_col).value
        if not name:
            continue
        seen.add(name)
        info = statuses.get(name)
        if info:
            value = 'Banked to Supabase: %s folders, %s columns (%s, %s)' % (
                info['folders'], info['columns'], info['formats'], info['date'])
            banked += 1
        else:
            value = NOT_BANKED
        cell = ws.cell(row=row, column=col, value=value)
        for attr in ('font', 'alignment', 'border'):
            setattr(cell, attr, copy(getattr(body, attr)))

    missing = sorted(set(statuses) - seen)
    wb.save(out_path)
    return banked, len(seen), missing


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('listing')
    ap.add_argument('statuses')
    ap.add_argument('-o', '--out')
    args = ap.parse_args()

    statuses = json.load(open(args.statuses, encoding='utf-8'))
    banked, rows, missing = refresh(args.listing, statuses, args.out or args.listing)

    sys.stderr.write('%d subject area rows, %d marked banked\n' % (rows, banked))
    for name in missing:
        # A banked subject area the register does not list is worth knowing about:
        # it means the name in otbi_meta and the name in the register disagree.
        sys.stderr.write('  WARNING: banked but not in the register: %s\n' % name)


if __name__ == '__main__':
    main()
